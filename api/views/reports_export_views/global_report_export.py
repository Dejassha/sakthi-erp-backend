import io
import json
import datetime
from datetime import timedelta
from django.utils import timezone
from django.http import HttpResponse
from rest_framework.decorators import api_view
from rest_framework.response import Response
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.db.models import F

from api.models import (
    product_material,
    product_details,
    programer_details,
    qa_details,
    qa_machine_details,
    acc_details,
    Quotation,
)
from api.views.utils import (
    apply_ag_grid_filter,
    parse_time_to_minutes,
    format_duration_as_time,
)


# ==============================================================================
# 1. Global Report Export (Selected / Filtered Rows)
# ==============================================================================
@api_view(["POST"])
def export_selected_rows(request):
    """
    Exports selected or filtered Global Report rows to Excel.
    Aligned with get_overall_details logic for field names, filtering, and totals.
    """
    try:
        headers = request.data.get("headers", [])
        rows = request.data.get("rows", [])
        filters_json = request.data.get("filters")
        include_totals = request.data.get("include_totals", True)

        # 1. FETCH DATA FROM DB IF ROWS NOT PROVIDED DIRECTLY
        if not rows:
            FIELD_MAP = {
                # Product
                "sheet_type": "product__sheet_type",
                "job_type": "product__job_type",
                "inward_slip_number": "product__inward_slip_number",
                "date": "product__date",
                "company_name": "product__company_name",
                "company": "product__company_name",
                "customer_name": "product__customer_name",
                "customer": "product__customer_name",
                "contact_no": "product__contact_no",
                "customer_dc_no": "product__customer_dc_no",
                "worker_no": "product__worker_no",
                "inward_created_by": "product__created_by",
                # Material
                "mat_type": "mat_type",
                "mat_grade": "mat_grade",
                "uid_no": "uid_no",
                "heat_no": "heat_no",
                "bay": "bay",
                "thick": "thick",
                "width": "width",
                "length": "length",
                "quantity": "quantity",
                "density": "density",
                "unit_weight": "unit_weight",
                "total_weight": "total_weight",
                "stock_due": "stock_due",
                "programer_status": "programer_status",
                "qa_status": "qa_status",
                "acc_status": "acc_status",
                "material_created_by": "created_by",
                # Programmer
                "program_no": "programer_details__program_no",
                "program_date": "programer_details__program_date",
                "processed_quantity": "programer_details__processed_quantity",
                "balance_quantity": "programer_details__balance_quantity",
                "processed_width": "programer_details__processed_width",
                "processed_length": "programer_details__processed_length",
                "used_weight": "programer_details__used_weight",
                "number_of_sheets": "programer_details__number_of_sheets",
                "cut_length_per_sheet": "programer_details__cut_length_per_sheet",
                "pierce_per_sheet": "programer_details__pierce_per_sheet",
                "processed_mins_per_sheet": "programer_details__processed_mins_per_sheet",
                "total_planned_hours": "programer_details__total_planned_hours",
                "total_no_of_sheets": "programer_details__total_no_of_sheets",
                "total_meters": "programer_details__total_meters",
                "total_piercing": "programer_details__total_piercing",
                "total_used_weight": "programer_details__total_used_weight",
                "programer_created_by": "programer_details__created_by",
                # QA
                "qa_processed_date": "qa_details__processed_date",
                "shift": "qa_details__shift",
                "qa_created_by": "qa_details__created_by",
                # Machine Logs
                "machine_name": "qa_details__qa_machine_details__machine_name",
                "machine_date": "qa_details__qa_machine_details__date",
                "start_time": "qa_details__qa_machine_details__start_time",
                "end_time": "qa_details__qa_machine_details__end_time",
                "machine_start": "qa_details__qa_machine_details__start_time",
                "machine_end": "qa_details__qa_machine_details__end_time",
                "machine_runtime": "qa_details__qa_machine_details__runtime",
                "machine_operator": "qa_details__qa_machine_details__operator",
                "machine_air": "qa_details__qa_machine_details__gas_type",
                "machine_gas_type": "qa_details__qa_machine_details__gas_type",
                # Accounts
                "invoice_no": "acc_details__invoice_no",
                "transporter_no": "acc_details__transporter_no",
                "payments_terms": "acc_details__payments_terms",
                "status": "acc_details__status",
                "acc_remarks": "acc_details__remarks",
                "acc_created_by": "acc_details__created_by",
            }

            material_qs = product_material.objects.select_related("product").distinct()

            filter_model = None
            if isinstance(filters_json, str):
                try:
                    filter_model = json.loads(filters_json)
                except:
                    pass
            elif isinstance(filters_json, dict):
                filter_model = filters_json

            if filter_model:
                for field, config in filter_model.items():
                    db_field = FIELD_MAP.get(field, field)
                    material_qs = apply_ag_grid_filter(material_qs, db_field, config)

            rows = []
            expanded_qs = material_qs.annotate(
                machine_log_id=F("qa_details__qa_machine_details__id")
            )
            paged_materials = (
                expanded_qs.prefetch_related(
                    "programer_details_set",
                    "qa_details_set",
                    "qa_details_set__qa_machine_details",
                    "acc_details_set",
                )
                .all()
                .order_by("-id")
            )

            for mat in paged_materials:
                prod = mat.product
                prog_list = list(mat.programer_details_set.all())
                prog = prog_list[0] if prog_list else None

                qa_list = list(mat.qa_details_set.all())
                qa = qa_list[0] if qa_list else None

                acc_list = list(mat.acc_details_set.all())
                acc = acc_list[0] if acc_list else None

                base_data = {
                    "material_id": mat.id,
                    "product_id": prod.id if prod else None,
                    "date": prod.date.isoformat() if prod and prod.date else "",
                    "job_type": prod.job_type if prod else "",
                    "inward_slip_number": prod.inward_slip_number if prod else "",
                    "customer": prod.customer_name if prod else "",
                    "customer_name": prod.customer_name if prod else "",
                    "company": prod.company_name if prod else "",
                    "company_name": prod.company_name if prod else "",
                    "sheet_type": prod.sheet_type if prod else "",
                    "worker_no": prod.worker_no if prod else "",
                    "customer_dc_no": prod.customer_dc_no if prod else "",
                    "contact_no": prod.contact_no if prod else "",
                    "inward_created_by": prod.created_by if prod else "",
                    "uid_no": mat.uid_no or "",
                    "heat_no": mat.heat_no or "",
                    "bay": mat.bay or "",
                    "mat_type": mat.mat_type or "",
                    "mat_grade": mat.mat_grade or "",
                    "thick": float(mat.thick) if mat.thick else 0,
                    "width": float(mat.width) if mat.width else 0,
                    "length": float(mat.length) if mat.length else 0,
                    "density": float(mat.density) if mat.density else 0,
                    "unit_weight": float(mat.unit_weight) if mat.unit_weight else 0,
                    "quantity": float(mat.quantity) if mat.quantity else 0,
                    "total_weight": float(mat.total_weight) if mat.total_weight else 0,
                    "stock_due": mat.stock_due or "",
                    "material_created_by": mat.created_by or "",
                    "programer_status": mat.programer_status or "",
                    "qa_status": mat.qa_status or "",
                    "acc_status": mat.acc_status or "",
                    "program_no": prog.program_no if prog else "",
                    "program_date": (
                        prog.program_date.isoformat() if prog and prog.program_date else ""
                    ),
                    "processed_quantity": (
                        float(prog.processed_quantity)
                        if prog and prog.processed_quantity
                        else 0
                    ),
                    "balance_quantity": (
                        float(prog.balance_quantity)
                        if prog and prog.balance_quantity
                        else 0
                    ),
                    "processed_width": (
                        float(prog.processed_width) if prog and prog.processed_width else 0
                    ),
                    "processed_length": (
                        float(prog.processed_length)
                        if prog and prog.processed_length
                        else 0
                    ),
                    "used_weight": (
                        float(prog.used_weight) if prog and prog.used_weight else 0
                    ),
                    "number_of_sheets": (
                        float(prog.number_of_sheets)
                        if prog and prog.number_of_sheets
                        else 0
                    ),
                    "cut_length_per_sheet": (
                        float(prog.cut_length_per_sheet)
                        if prog and prog.cut_length_per_sheet
                        else 0
                    ),
                    "pierce_per_sheet": (
                        float(prog.pierce_per_sheet)
                        if prog and prog.pierce_per_sheet
                        else 0
                    ),
                    "processed_mins_per_sheet": (
                        float(prog.processed_mins_per_sheet)
                        if prog and prog.processed_mins_per_sheet
                        else 0
                    ),
                    "total_planned_hours": (
                        format_duration_as_time(prog.total_planned_hours)
                        if prog and prog.total_planned_hours
                        else ""
                    ),
                    "total_meters": (
                        float(prog.total_meters) if prog and prog.total_meters else 0
                    ),
                    "total_piercing": (
                        float(prog.total_piercing) if prog and prog.total_piercing else 0
                    ),
                    "total_used_weight": (
                        float(prog.total_used_weight)
                        if prog and prog.total_used_weight
                        else 0
                    ),
                    "total_no_of_sheets": (
                        float(prog.total_no_of_sheets)
                        if prog and prog.total_no_of_sheets
                        else 0
                    ),
                    "programer_created_by": prog.created_by if prog else "",
                    "qa_processed_date": (
                        qa.processed_date.isoformat() if qa and qa.processed_date else ""
                    ),
                    "shift": qa.shift if qa else "",
                    "qa_created_by": qa.created_by if qa else "",
                    "invoice_no": acc.invoice_no if acc else "",
                    "transporter_no": acc.transporter_no if acc else "",
                    "payments_terms": acc.payments_terms if acc else "",
                    "status": acc.status if acc else "",
                    "acc_remarks": acc.remarks if acc else "",
                    "acc_created_by": acc.created_by if acc else "",
                }

                machine_log_id = getattr(mat, "machine_log_id", None)
                logs = qa.qa_machine_details.all() if qa else []
                row = base_data.copy()

                if machine_log_id and logs:
                    log = next((l for l in logs if l.id == machine_log_id), None)
                    if log:
                        row.update(
                            {
                                "machine_name": log.machine_name or "",
                                "machine_date": log.date.isoformat() if log.date else "",
                                "machine_start": format_duration_as_time(log.start_time),
                                "machine_end": format_duration_as_time(log.end_time),
                                "machine_runtime": format_duration_as_time(log.runtime)
                                or "00:00",
                                "machine_air": log.gas_type or "",
                                "machine_gas_type": log.gas_type or "",
                                "machine_operator": log.operator or "",
                            }
                        )
                    else:
                        row.update(
                            {
                                "machine_name": "",
                                "machine_date": "",
                                "machine_start": "",
                                "machine_end": "",
                                "machine_runtime": "",
                                "machine_operator": "",
                                "machine_air": "",
                                "machine_gas_type": "",
                            }
                        )
                else:
                    row.update(
                        {
                            "machine_name": "",
                            "machine_date": "",
                            "machine_start": "",
                            "machine_end": "",
                            "machine_runtime": "",
                            "machine_operator": "",
                            "machine_air": "",
                            "machine_gas_type": "",
                        }
                    )
                rows.append(row)

        if not rows or not headers:
            return Response({"detail": "No data provided"}, status=400)

        wb = Workbook()
        ws = wb.active
        ws.title = "Global Report"

        header_labels = [
            str(h["header"]) if isinstance(h, dict) and h.get("header") is not None else str(h.get("label", ""))
            for h in headers
        ]
        fields = [
            h["field"] if isinstance(h, dict) and "field" in h else h.get("value", "")
            for h in headers
        ]

        total_cols = len(header_labels)
        last_col_letter = get_column_letter(max(total_cols, 1))

        # ---------------- Styles ----------------
        title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        title_fill = PatternFill("solid", fgColor="0F172A")
        title_align = Alignment(horizontal="center", vertical="center")

        subtitle_font = Font(name="Calibri", size=9, italic=True, color="94A3B8")
        subtitle_fill = PatternFill("solid", fgColor="0F172A")
        subtitle_align = Alignment(horizontal="center", vertical="center")

        header_fill = PatternFill("solid", fgColor="2563EB")  # Tailwind blue-600
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        center_align = Alignment(
            horizontal="center", vertical="center", wrap_text=True, indent=1
        )

        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        total_fill = PatternFill("solid", fgColor="E5E7EB")
        total_font = Font(bold=True)

        now_local = (
            timezone.localtime(timezone.now())
            if timezone.is_aware(timezone.now())
            else datetime.datetime.now()
        )
        now_str = now_local.strftime("%d-%b-%Y %I:%M %p")

        report_title_custom = request.data.get("report_title")
        title_text = (
            f"SAKTHI LASER TECHNOLOGY - {str(report_title_custom).upper()}"
            if report_title_custom
            else "SAKTHI LASER TECHNOLOGY - GLOBAL REPORT"
        )

        # ---------------- Title Banner (Rows 1 & 2) ----------------
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = title_align
        ws.row_dimensions[1].height = 24

        for c_idx in range(1, total_cols + 1):
            c = ws.cell(row=1, column=c_idx)
            c.fill = title_fill
            c.font = title_font

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        sub_cell = ws.cell(row=2, column=1, value=f"Generated on {now_str}")
        sub_cell.font = subtitle_font
        sub_cell.fill = subtitle_fill
        sub_cell.alignment = subtitle_align
        ws.row_dimensions[2].height = 16

        for c_idx in range(1, total_cols + 1):
            c = ws.cell(row=2, column=c_idx)
            c.fill = subtitle_fill
            c.font = subtitle_font

        # Row 3: Blank spacing row
        ws.row_dimensions[3].height = 6

        # ---------------- Column Headers (Row 4) ----------------
        for col_idx, label in enumerate(header_labels, 1):
            cell = ws.cell(row=4, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border
        ws.row_dimensions[4].height = 22

        # ---------------- Data Rows (Row 5 onwards) ----------------
        for row in rows:
            ws.append([row.get(f, "") for f in fields])

        sum_columns = [
            "thick",
            "width",
            "length",
            "unit_weight",
            "quantity",
            "total_weight",
            "processed_quantity",
            "balance_quantity",
            "processed_width",
            "processed_length",
            "used_weight",
            "number_of_sheets",
            "cut_length_per_sheet",
            "pierce_per_sheet",
            "processed_mins_per_sheet",
            "total_meters",
            "total_piercing",
            "total_used_weight",
            "total_no_of_sheets",
            "no_of_sheets",
        ]
        time_sum_columns = [
            "total_planned_hours",
            "machine_runtime",
        ]

        # ---------------- Totals Calculation ----------------
        if include_totals and len(rows) >= 1:
            unique_mat_rows = {}
            for row in rows:
                mid = row.get("material_id")
                if mid and mid not in unique_mat_rows:
                    unique_mat_rows[mid] = row
                elif not mid:
                    unique_mat_rows[id(row)] = row

            totals = {f: 0 for f in fields}

            for row in unique_mat_rows.values():
                for f in fields:
                    if f in sum_columns:
                        try:
                            val = row.get(f)
                            if val is not None and str(val).strip() != "":
                                totals[f] += float(val)
                        except (ValueError, TypeError):
                            pass
                    elif f in time_sum_columns:
                        val = row.get(f)
                        if val is not None and str(val).strip() != "":
                            totals[f] += parse_time_to_minutes(val)

            if "machine_runtime" in fields:
                totals["machine_runtime"] = sum(
                    parse_time_to_minutes(r.get("machine_runtime", 0)) for r in rows
                )

            total_row_data = []
            for f in fields:
                if f == fields[0]:
                    total_row_data.append("TOTAL")
                elif f in sum_columns:
                    val = totals[f]
                    # Round float sums to 3 decimals to avoid binary floating point precision artifacts
                    total_row_data.append(round(val, 3) if isinstance(val, float) else val)
                elif f in time_sum_columns:
                    mins = totals[f]
                    total_row_data.append(
                        format_duration_as_time(timedelta(minutes=mins))
                    )
                else:
                    total_row_data.append("")

            ws.append(total_row_data)

        last_row = ws.max_row

        # ---------------- TOTAL row styling ----------------
        if include_totals and len(rows) >= 1:
            for col in range(1, len(fields) + 1):
                cell = ws.cell(row=last_row, column=col)
                cell.fill = total_fill
                cell.font = total_font
                cell.border = border

        # ---------------- Borders + Alignment ----------------
        left_align_fields = {
            "company_name",
            "customer_name",
            "company",
            "customer",
            "worker_no",
            "acc_remarks",
            "remarks",
        }
        right_align_fields = set(sum_columns) | set(time_sum_columns) | {
            "stock_due",
            "density",
            "cut_length_per_sheet",
            "pierce_per_sheet",
            "number_of_sheets",
        }

        for row in ws.iter_rows(min_row=5, max_row=last_row):
            for col_idx, cell in enumerate(row):
                field_name = fields[col_idx] if col_idx < len(fields) else ""
                cell.border = border
                if field_name in right_align_fields:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif field_name in left_align_fields:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # ---------------- Auto Column Width ----------------
        for i, header in enumerate(header_labels, 1):
            max_len = len(str(header)) if header else 0
            col_letter = get_column_letter(i)
            for cell in ws[col_letter]:
                if cell.row >= 4 and cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

        # ---------------- Freeze Header ----------------
        ws.freeze_panes = "A5"

        # ---------------- Save ----------------
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        response = HttpResponse(
            stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="production_export.xlsx"'

        return response

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({"error": str(e)}, status=500)
